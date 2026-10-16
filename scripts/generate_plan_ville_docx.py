# -*- coding: utf-8 -*-
"""Genera plan CostaBIM ES + FI (terminología finlandesa de määrälaskenta)."""
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import os
from datetime import date

DOWNLOADS = os.path.join(os.environ["USERPROFILE"], "Downloads")

BLUE = "1F4E79"
BLUE_LIGHT = "D6E3F0"
GREEN = "548235"
GREEN_LIGHT = "C6EFCE"
YELLOW = "BF8F00"
YELLOW_LIGHT = "FFEB9C"
RED = "C00000"
RED_LIGHT = "FFC7CE"
GRAY = "7F7F7F"
GRAY_LIGHT = "F2F2F2"
WHITE = "FFFFFF"
PAST_HEADER = "2E75B6"
FUTURE_HEADER = "548235"


def set_cell_shading(cell, fill_hex: str):
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill_hex)
    shading.set(qn("w:val"), "clear")
    cell._tc.get_or_add_tcPr().append(shading)


def set_cell_borders(cell, color="A9A9A9", sz="4"):
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), sz)
        el.set(qn("w:color"), color)
        tcBorders.append(el)
    tcPr.append(tcBorders)


def set_cell_text(
    cell,
    text,
    *,
    bold=False,
    size=10,
    color=None,
    align=WD_ALIGN_PARAGRAPH.LEFT,
):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = align
    run = p.add_run(str(text))
    run.bold = bold
    run.font.size = Pt(size)
    run.font.name = "Calibri"
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_heading(doc, text, level=1):
    return doc.add_heading(text, level=level)


def add_para(doc, text, bold=False, size=11, italic=False):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = bold
    run.italic = italic
    run.font.size = Pt(size)
    run.font.name = "Calibri"
    return p


def add_table(doc, headers, rows, col_widths=None, header_fill=BLUE):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(headers):
        set_cell_text(
            table.rows[0].cells[i],
            h,
            bold=True,
            size=10,
            color=WHITE,
            align=WD_ALIGN_PARAGRAPH.CENTER,
        )
        set_cell_shading(table.rows[0].cells[i], header_fill)
    for ri, row in enumerate(rows):
        for ci, val in enumerate(row):
            set_cell_text(table.rows[ri + 1].cells[ci], val, size=10)
            if ri % 2 == 1:
                set_cell_shading(table.rows[ri + 1].cells[ci], GRAY_LIGHT)
    if col_widths:
        for i, w in enumerate(col_widths):
            for row in table.rows:
                row.cells[i].width = Cm(w)
    doc.add_paragraph()
    return table


def add_flow_diagram(doc, title, steps):
    add_para(doc, title, bold=True, size=11)
    n = len(steps)
    table = doc.add_table(rows=1, cols=n * 2 - 1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    col = 0
    for i, step in enumerate(steps):
        c = table.rows[0].cells[col]
        set_cell_text(c, step, bold=True, size=9, align=WD_ALIGN_PARAGRAPH.CENTER)
        set_cell_shading(c, BLUE_LIGHT)
        set_cell_borders(c, BLUE, "8")
        if i < n - 1:
            arr = table.rows[0].cells[col + 1]
            set_cell_text(
                arr, "→", bold=True, size=14, align=WD_ALIGN_PARAGRAPH.CENTER, color=BLUE
            )
            col += 2
        else:
            col += 1
    doc.add_paragraph()


def add_cronograma(doc, title, periods, rows, header_fill):
    add_para(doc, title, bold=True, size=11)
    cols = 1 + len(periods)
    table = doc.add_table(rows=1 + len(rows), cols=cols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    set_cell_text(
        table.rows[0].cells[0],
        "",
        bold=True,
        size=9,
        color=WHITE,
        align=WD_ALIGN_PARAGRAPH.CENTER,
    )
    set_cell_shading(table.rows[0].cells[0], header_fill)
    for i, p in enumerate(periods):
        set_cell_text(
            table.rows[0].cells[i + 1],
            p,
            bold=True,
            size=8,
            color=WHITE,
            align=WD_ALIGN_PARAGRAPH.CENTER,
        )
        set_cell_shading(table.rows[0].cells[i + 1], header_fill)

    fill_map = {
        "done": GREEN_LIGHT,
        "partial": YELLOW_LIGHT,
        "todo": RED_LIGHT,
        "focus": BLUE_LIGHT,
        "empty": WHITE,
    }
    label_map = {
        "done": "●",
        "partial": "◐",
        "todo": "○",
        "focus": "◆",
        "empty": "",
    }

    for ri, (actividad, estados) in enumerate(rows):
        set_cell_text(table.rows[ri + 1].cells[0], actividad, bold=True, size=9)
        set_cell_shading(table.rows[ri + 1].cells[0], GRAY_LIGHT)
        for ci, est in enumerate(estados):
            cell = table.rows[ri + 1].cells[ci + 1]
            set_cell_text(
                cell, label_map.get(est, ""), bold=True, size=12, align=WD_ALIGN_PARAGRAPH.CENTER
            )
            set_cell_shading(cell, fill_map.get(est, WHITE))
            set_cell_borders(cell, "BFBFBF", "4")

    table.rows[0].cells[0].width = Cm(5.5)
    for i in range(len(periods)):
        for row in table.rows:
            row.cells[i + 1].width = Cm(2.0)

    return table


def set_first_crono_header(table, label):
    set_cell_text(
        table.rows[0].cells[0],
        label,
        bold=True,
        size=9,
        color=WHITE,
        align=WD_ALIGN_PARAGRAPH.CENTER,
    )


def add_legend(doc, items):
    leg = doc.add_paragraph()
    for text, col in items:
        run = leg.add_run(text)
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor.from_string(col)
        run.font.name = "Calibri"
    doc.add_paragraph()


def fill_map_safe(badge_hex):
    return {
        GREEN: GREEN_LIGHT,
        YELLOW: YELLOW_LIGHT,
        RED: RED_LIGHT,
    }.get(badge_hex, GRAY_LIGHT)


def add_pipeline(doc, stages, compact_labels, compact_colors, compact_status, headers):
    add_para(doc, headers["read"], italic=True, size=10)

    table = doc.add_table(rows=1 + len(stages), cols=5)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdrs = headers["cols"]
    widths = [1.0, 4.2, 2.4, 4.5, 4.5]
    for i, h in enumerate(hdrs):
        set_cell_text(
            table.rows[0].cells[i],
            h,
            bold=True,
            size=9,
            color=WHITE,
            align=WD_ALIGN_PARAGRAPH.CENTER,
        )
        set_cell_shading(table.rows[0].cells[i], BLUE)
        table.rows[0].cells[i].width = Cm(widths[i])

    for ri, s in enumerate(stages):
        row = table.rows[ri + 1]
        set_cell_text(
            row.cells[0],
            s["n"],
            bold=True,
            size=14,
            align=WD_ALIGN_PARAGRAPH.CENTER,
            color=s["badge"],
        )
        set_cell_shading(row.cells[0], s["fill"])
        set_cell_borders(row.cells[0], s["badge"], "12")
        set_cell_text(row.cells[1], s["title"], bold=True, size=10)
        set_cell_shading(row.cells[1], s["fill"])
        set_cell_text(
            row.cells[2],
            s["status"],
            bold=True,
            size=10,
            color=WHITE,
            align=WD_ALIGN_PARAGRAPH.CENTER,
        )
        set_cell_shading(row.cells[2], s["badge"])
        set_cell_text(row.cells[3], s["detail"], size=9)
        set_cell_shading(row.cells[3], WHITE)
        set_cell_text(row.cells[4], s["blocker"], size=9)
        set_cell_shading(
            row.cells[4], GRAY_LIGHT if s["blocker"] not in ("—", "–") else WHITE
        )
        for ci in range(5):
            row.cells[ci].width = Cm(widths[ci])

    doc.add_paragraph()
    add_para(doc, headers["compact"], bold=True, size=10)
    pipe = doc.add_table(rows=2, cols=11)
    pipe.alignment = WD_TABLE_ALIGNMENT.CENTER
    col = 0
    for i in range(6):
        c0 = pipe.rows[0].cells[col]
        set_cell_text(
            c0,
            compact_labels[i],
            bold=True,
            size=8,
            color=WHITE,
            align=WD_ALIGN_PARAGRAPH.CENTER,
        )
        set_cell_shading(c0, compact_colors[i])
        set_cell_borders(c0, compact_colors[i], "10")
        c1 = pipe.rows[1].cells[col]
        set_cell_text(
            c1, compact_status[i], bold=True, size=11, align=WD_ALIGN_PARAGRAPH.CENTER
        )
        set_cell_shading(c1, fill_map_safe(compact_colors[i]))
        if i < 5:
            a0 = pipe.rows[0].cells[col + 1]
            set_cell_text(
                a0, "→", bold=True, size=12, align=WD_ALIGN_PARAGRAPH.CENTER, color=GRAY
            )
            set_cell_text(pipe.rows[1].cells[col + 1], "", size=8)
            col += 2
        else:
            col += 1
    doc.add_paragraph()


# ---------------------------------------------------------------------------
# Contenido ES (revisado + términos FI que usa Costa)
# ---------------------------------------------------------------------------

ES = {
    "out": "CostaBIM_Plan_Ville_ES_v3.docx",
    "lang": "es",
    "cover_sub": "Plan de producto: del modelo IFC al määräluettelo (lista de cantidades)",
    "cover_note": "Borrador v3 ES — contraste con versión FI. Producto: CostaBIM.",
    "h1": "1. Resumen",
    "intro": "Este documento responde a cuatro preguntas:",
    "q_headers": ["Pregunta", "Respuesta breve"],
    "q_rows": [
        [
            "¿Cuál es el objetivo?",
            "Sacar määrät (cantidades) fiables del IFC a Measurement Data y al määräluettelo de Costa.",
        ],
        [
            "¿Qué NO es el foco de este plan?",
            "Hinnoittelu / facturación con IA: hilo distinto (hay avances en la empresa). Este plan se centra en määrälaskenta.",
        ],
        [
            "¿Dónde estamos?",
            "CostaBIM: visor + Excel operativos; falta Talo 2000 -koodi en muchos IFC y reglas Littera cerradas.",
        ],
        [
            "¿Cuándo producción (tuotantokäyttö)?",
            "Ver cronograma §6 — piloto + validación.",
        ],
    ],
    "h2": "2. Contexto para määrälaskenta (Finlandia / Costa)",
    "ctx": (
        "En Costa, una medición útil (mittaus) no es solo “geometría del IFC”. "
        "El equipo trabaja con un objeto medible (mittauskohde) de tres datos:"
    ),
    "field_headers": ["Campo", "Ejemplo", "Para qué sirve"],
    "field_rows": [
        ["IFC Type", "IfcSlab, IfcWindow", "Tipo de elemento (rakennusosa en IFC)"],
        ["Name", "Floor:O32", "Variante / nombre Revit normalizado"],
        [
            "TALO Code",
            "1.2.3.5",
            "Talo 2000 Hankenimikkeistö (nimike / clasificación de obra)",
        ],
    ],
    "talo_para": (
        "Talo 2000 Hankenimikkeistö es el vehículo de clasificación. "
        "Sobre Type + Name + TALO se calculan määrät con unidades finlandesas "
        "(m², m³, m, kpl, kg y JM = juoksumetri) y, cuando aplica, códigos Littera "
        "del catálogo Costa (que alimentan el määräluettelo)."
    ),
    "schema1": "Esquema 1 — Mittauskohde Costa (objeto medible)",
    "flow1_title": "Flujo lógico (määrälaskenta):",
    "flow1": [
        "Modelo IFC\n(Revit)",
        "Talo 2000\nHankenimikkeistö",
        "Määrät\n(cantidades)",
        "Määräluettelo\n(Measurement Data)",
    ],
    "flow1_note": (
        "El tramo crítico: määrät → määräluettelo. "
        "TALO y Littera hacen ese tramo utilizable y repetible en la práctica finlandesa."
    ),
    "h3": "3. Dos formas de asignar TALO (nimike) al modelo",
    "via_headers": ["Vía", "Cuándo se usa", "Responsable negocio", "Estado"],
    "via_rows": [
        [
            "Plugin Revit (interno)",
            "Proyectos modelados en Costa",
            "Topi",
            "Desarrollado — adopción interna por definir",
        ],
        [
            "Claude + CostaBIM (MCP)",
            "IFC de clientes sin Talo-koodi",
            "Equipo + revisión",
            "Operativo en pruebas — requiere revisión humana (tarkistus)",
        ],
    ],
    "flow2_title": "Esquema 2 — Cliente envía IFC sin TALO:",
    "flow2": [
        "IFC cliente",
        "CostaBIM\n+ Claude",
        "Escaneo y\npropuesta TALO",
        "IFC con nimike\n(o store CostaBIM)",
        "Excel\nMeasurement Data\n→ Määräluettelo",
    ],
    "h4": "4. Qué se ha hecho (~8 meses) — cronograma",
    "crono_a_title": "Cronograma A — Trabajo realizado (aprox. últimos 8 meses)",
    "crono_a_first": "Actividad / bloque",
    "crono_a_periods": ["Mes 1–2", "Mes 3–4", "Mes 5–6", "Mes 7–8", "Hoy"],
    "crono_a_rows": [
        ("Visor IFC (CostaBIM)", ["focus", "done", "done", "done", "done"]),
        ("Extracción de määrät", ["partial", "focus", "done", "done", "done"]),
        ("Cost Analysis + Excel Measurement Data", ["empty", "partial", "focus", "done", "done"]),
        ("TALO — plugin Revit (interno)", ["empty", "partial", "focus", "done", "done"]),
        ("TALO — Claude / MCP (IFC cliente)", ["empty", "empty", "partial", "focus", "partial"]),
        ("Littera (reglas + casos especiales)", ["empty", "empty", "partial", "partial", "partial"]),
        ("Validación en proyectos reales", ["empty", "empty", "empty", "partial", "partial"]),
    ],
    "legend": [
        ("● Hecho / listo  ", GREEN),
        ("◐ Parcial / en curso  ", YELLOW),
        ("○ Pendiente  ", RED),
        ("◆ Enfoque del periodo  ", BLUE),
        ("(vacío = no aplica)", GRAY),
    ],
    "summary4": (
        "Resumen: la tubería técnica de CostaBIM (visor → määrät → Excel) está operativa; "
        "lo que frena el uso diario es Talo-luokittelu incompleta y reglas Littera aún abiertas."
    ),
    "h5": "5. Brechas actuales (honestas)",
    "gap_headers": ["Brecha", "Impacto", "Qué se necesita"],
    "gap_rows": [
        [
            "IFC sin Talo-koodi",
            "Mittauskohde incompleto → Littera / kustannuslaskenta débil",
            "Re-escaneo + asignación TALO (§3)",
        ],
        [
            "Reglas Littera no cerradas por tipo/TALO",
            "Excel no refleja el catálogo Costa (→ määräluettelo) de forma uniforme",
            "Definición negocio (Jesse) + validación Työkirja",
        ],
        [
            "Validación en proyectos reales limitada",
            "Aún no hay tuotantokäyttö",
            "Proyecto piloto + criterios de aceptación",
        ],
        [
            "Prioridad de desarrollo dirigido de CostaBIM",
            "El trabajo diario de medición suele ir por delante del desarrollo de la herramienta → avance más lento",
            "Decisión conjunta de foco de desarrollo y bloques regulares de avance",
        ],
    ],
    "h_schema3": "Esquema 3 — Estado del tubo principal (lectura clave)",
    "pipe_headers": {
        "read": (
            "Lectura rápida: verde = listo; amarillo = funciona pero incompleto; "
            "rojo = bloquea tuotantokäyttö."
        ),
        "cols": ["#", "Etapa del tubo", "Semáforo", "Qué significa", "Qué falta / riesgo"],
        "compact": "Vista compacta del tubo (izq. → der.):",
    },
    "stages": [
        {
            "n": "1",
            "title": "Leer IFC en CostaBIM",
            "status": "LISTO",
            "fill": GREEN_LIGHT,
            "badge": GREEN,
            "detail": "CostaBIM carga y muestra el modelo 3D.",
            "blocker": "—",
        },
        {
            "n": "2",
            "title": "Extraer määrät",
            "status": "LISTO",
            "fill": GREEN_LIGHT,
            "badge": GREEN,
            "detail": "m², m³, m, kg, kpl, JM por tipo IFC.",
            "blocker": "Validación fina por tipo (en curso).",
        },
        {
            "n": "3",
            "title": "TALO en el elemento",
            "status": "PARCIAL",
            "fill": YELLOW_LIGHT,
            "badge": YELLOW,
            "detail": "Mittauskohde = Type + Name + TALO.",
            "blocker": "Muchos IFC de cliente llegan sin Talo-koodi.",
        },
        {
            "n": "4",
            "title": "Littera alineado a Costa",
            "status": "PARCIAL",
            "fill": YELLOW_LIGHT,
            "badge": YELLOW,
            "detail": "Littera-koodit desde Type/Name/TALO.",
            "blocker": "Reglas de negocio por cerrar (Jesse).",
        },
        {
            "n": "5",
            "title": "Excel Measurement Data → Määräluettelo",
            "status": "LISTO",
            "fill": GREEN_LIGHT,
            "badge": GREEN,
            "detail": "Export / import desde Cost Analysis.",
            "blocker": "Calidad depende de etapas 3–4.",
        },
        {
            "n": "6",
            "title": "Tuotantokäyttö (uso regular)",
            "status": "PENDIENTE",
            "fill": RED_LIGHT,
            "badge": RED,
            "detail": "Equipo de määrälaskenta sin soporte diario de desarrollo.",
            "blocker": "Piloto + criterios de aceptación.",
        },
    ],
    "compact_labels": ["IFC", "Määrät", "TALO", "Littera", "Excel", "Prod."],
    "h6": "6. Plan futuro y calendario (orientativo)",
    "crono_b_title": "Cronograma B — Plan hacia tuotantokäyttö (periodos a confirmar)",
    "crono_b_first": "Actividad / bloque",
    "crono_b_periods": ["F1\nSem 1–4", "F2\nSem 4–8", "F3\nSem 8–12", "F4\nMes 3", "F5\nProd."],
    "crono_b_rows": [
        ("Määrät validados (tipos MVP)", ["focus", "done", "done", "done", "done"]),
        ("TALO cubierto (Revit + MCP clientes)", ["partial", "focus", "done", "done", "done"]),
        ("Reglas Littera (TALO → Littera)", ["empty", "partial", "focus", "done", "done"]),
        ("Piloto usuarios Costa (sin soporte diario)", ["empty", "empty", "partial", "focus", "done"]),
        ("Producción MVP (tuotantokäyttö)", ["empty", "empty", "empty", "partial", "focus"]),
    ],
    "phase_headers": ["Fase", "Objetivo", "Entregable", "Periodo (completar)"],
    "phase_rows": [
        ["F1", "Määrät validados por tipos IFC acordados", "Excel referencia aceptado internamente", "[Sem 1–4]"],
        ["F2", "TALO cubierto (Revit interno + MCP clientes)", "Procedimiento + 1 piloto", "[Sem 4–8]"],
        ["F3", "Reglas Littera (TALO → Littera)", "Documento reglas + piloto", "[Sem 8–12]"],
        ["F4", "Piloto con usuarios Costa (sin soporte diario)", "Go / no-go", "[Mes 3]"],
        ["F5", "Producción MVP (tuotantokäyttö)", "Fecha producción", "[Completar Jesse]"],
    ],
    "end": "— Fin borrador v3 ES —",
}


# ---------------------------------------------------------------------------
# Contenido FI
# ---------------------------------------------------------------------------

FI = {
    "out": "CostaBIM_Plan_Ville_FI_v3.docx",
    "lang": "fi",
    "cover_sub": "Tuotesuunnitelma: IFC-mallista määräluetteloon",
    "cover_note": "Luonnos v3 FI — vertailu ES-versioon. Tuote: CostaBIM.",
    "h1": "1. Tiivistelmä",
    "intro": "Tämä asiakirja vastaa neljään kysymykseen:",
    "q_headers": ["Kysymys", "Lyhyt vastaus"],
    "q_rows": [
        [
            "Mikä on tavoite?",
            "Luotettavat määrät IFC-mallista CostaBIM Measurement Dataan ja määräluetteloon.",
        ],
        [
            "Mikä EI ole tämän suunnitelman fokus?",
            "Hinnoittelu / laskutus tekoälyllä: erillinen kokonaisuus (yrityksessä on jo edistystä). Tämä suunnitelma keskittyy määrälaskentaan.",
        ],
        [
            "Missä olemme?",
            "CostaBIM: katselin + Excel toimivat; monista IFC-malleista puuttuu Talo-koodi, Littera-säännöt kesken.",
        ],
        [
            "Milloin tuotantokäyttö?",
            "Ks. aikataulu §6 — pilotti + validointi.",
        ],
    ],
    "h2": "2. Määrälaskennan konteksti (Suomi / Costa)",
    "ctx": (
        "Costassa käyttökelpoinen mittaus ei ole vain “IFC-geometriaa”. "
        "Tiimi käyttää kolmen tiedon mittauskohdetta:"
    ),
    "field_headers": ["Kenttä", "Esimerkki", "Mihin käytetään"],
    "field_rows": [
        ["IFC Type", "IfcSlab, IfcWindow", "Elementin tyyppi (rakennusosa IFC:ssä)"],
        ["Name", "Floor:O32", "Variantti / normalisoitu Revit-nimi"],
        [
            "TALO Code",
            "1.2.3.5",
            "Talo 2000 Hankenimikkeistö (hankkeen nimike)",
        ],
    ],
    "talo_para": (
        "Talo 2000 Hankenimikkeistö on luokittelun perusta. "
        "Type + Name + TALO -yhdistelmästä lasketaan määrät suomalaisilla yksiköillä "
        "(m², m³, m, kpl, kg ja JM = juoksumetri) ja tarvittaessa Costan Littera-koodit, "
        "jotka syöttävät määräluetteloa."
    ),
    "schema1": "Kaavio 1 — Costan mittauskohde",
    "flow1_title": "Looginen kulku (määrälaskenta):",
    "flow1": [
        "IFC-malli\n(Revit)",
        "Talo 2000\nHankenimikkeistö",
        "Määrät",
        "Määräluettelo\n(Measurement Data)",
    ],
    "flow1_note": (
        "Kriittinen osuus: määrät → määräluettelo. "
        "TALO ja Littera tekevät tästä toistettavan suomalaisessa määrälaskentakäytännössä."
    ),
    "h3": "3. Kaksi tapaa liittää TALO (nimike) malliin",
    "via_headers": ["Tapa", "Milloin käytetään", "Vastuu (liiketoiminta)", "Tila"],
    "via_rows": [
        [
            "Revit-lisäosa (sisäinen)",
            "Costan omat mallinnusprojektit",
            "Topi",
            "Kehitetty — sisäinen käyttöönotto auki",
        ],
        [
            "Claude + CostaBIM (MCP)",
            "Asiakkaan IFC ilman Talo-koodia",
            "Tiimi + tarkistus",
            "Koe käytössä — vaatii ihmistarkistuksen",
        ],
    ],
    "flow2_title": "Kaavio 2 — Asiakas lähettää IFC:n ilman TALOa:",
    "flow2": [
        "Asiakkaan IFC",
        "CostaBIM\n+ Claude",
        "Skannaus &\nTALO-ehdotus",
        "IFC nimikkeillä\n(tai CostaBIM-store)",
        "Excel\nMeasurement Data\n→ Määräluettelo",
    ],
    "h4": "4. Mitä on tehty (~8 kk) — aikataulu",
    "crono_a_title": "Aikataulu A — Tehty työ (n. viimeiset 8 kk)",
    "crono_a_first": "Toiminto / kokonaisuus",
    "crono_a_periods": ["Kk 1–2", "Kk 3–4", "Kk 5–6", "Kk 7–8", "Nyt"],
    "crono_a_rows": [
        ("IFC-katselin (CostaBIM)", ["focus", "done", "done", "done", "done"]),
        ("Määrien laskenta (määrät)", ["partial", "focus", "done", "done", "done"]),
        ("Cost Analysis + Excel Measurement Data", ["empty", "partial", "focus", "done", "done"]),
        ("TALO — Revit-lisäosa (sisäinen)", ["empty", "partial", "focus", "done", "done"]),
        ("TALO — Claude / MCP (asiakas-IFC)", ["empty", "empty", "partial", "focus", "partial"]),
        ("Littera (säännöt + erikoistapaukset)", ["empty", "empty", "partial", "partial", "partial"]),
        ("Validointi oikeissa hankkeissa", ["empty", "empty", "empty", "partial", "partial"]),
    ],
    "legend": [
        ("● Valmis  ", GREEN),
        ("◐ Osittainen / käynnissä  ", YELLOW),
        ("○ Avoin  ", RED),
        ("◆ Jakson painopiste  ", BLUE),
        ("(tyhjä = ei koske)", GRAY),
    ],
    "summary4": (
        "Yhteenveto: CostaBIMin tekninen putki (katselin → määrät → Excel) toimii; "
        "päivittäistä käyttöä hidastavat puutteellinen Talo-luokittelu ja avoimet Littera-säännöt."
    ),
    "h5": "5. Nykyiset aukot (rehellisesti)",
    "gap_headers": ["Aukko", "Vaikutus", "Mitä tarvitaan"],
    "gap_rows": [
        [
            "IFC ilman Talo-koodia",
            "Mittauskohde vajaana → Littera / kustannuslaskenta heikko",
            "Uudelleenskannaus + TALO-kohdistus (§3)",
        ],
        [
            "Littera-säännöt auki tyyppi/TALO -tasolla",
            "Excel ei vastaa Costan katalogia tasaisesti (→ määräluettelo)",
            "Liiketoimintasäännöt (Jesse) + Työkirja-validointi",
        ],
        [
            "Rajallinen validointi oikeissa hankkeissa",
            "Ei vielä tuotantokäyttöä",
            "Pilottihanke + hyväksymiskriteerit",
        ],
        [
            "CostaBIMin ohjatun kehityksen painoarvo",
            "Arjen määrälaskenta priorisoituu usein työkalun kehityksen edelle → eteneminen hidasta",
            "Yhteinen päätös kehityspainopisteestä ja säännöllisistä kehitysjaksoista",
        ],
    ],
    "h_schema3": "Kaavio 3 — Pääputken tila (tärkein lukema)",
    "pipe_headers": {
        "read": (
            "Pikaluku: vihreä = valmis; keltainen = toimii mutta vajaana; "
            "punainen = estää tuotantokäytön."
        ),
        "cols": ["#", "Putken vaihe", "Liikennevalo", "Mitä tarkoittaa", "Puute / riski"],
        "compact": "Kompakti putkinäkymä (vasemmalta oikealle):",
    },
    "stages": [
        {
            "n": "1",
            "title": "IFC CostaBIMissä",
            "status": "VALMIS",
            "fill": GREEN_LIGHT,
            "badge": GREEN,
            "detail": "CostaBIM lataa ja näyttää 3D-mallin.",
            "blocker": "—",
        },
        {
            "n": "2",
            "title": "Määrien laskenta",
            "status": "VALMIS",
            "fill": GREEN_LIGHT,
            "badge": GREEN,
            "detail": "m², m³, m, kg, kpl, JM IFC-tyypeittäin.",
            "blocker": "Hienosäätö tyyppikohtaisesti (käynnissä).",
        },
        {
            "n": "3",
            "title": "TALO elementissä",
            "status": "OSITTAINEN",
            "fill": YELLOW_LIGHT,
            "badge": YELLOW,
            "detail": "Mittauskohde = Type + Name + TALO.",
            "blocker": "Monissa asiakas-IFC:issä ei ole Talo-koodia.",
        },
        {
            "n": "4",
            "title": "Littera Costan mukaiseksi",
            "status": "OSITTAINEN",
            "fill": YELLOW_LIGHT,
            "badge": YELLOW,
            "detail": "Littera-koodit Type/Name/TALO -pohjalta.",
            "blocker": "Liiketoimintasäännöt auki (Jesse).",
        },
        {
            "n": "5",
            "title": "Excel Measurement Data → Määräluettelo",
            "status": "VALMIS",
            "fill": GREEN_LIGHT,
            "badge": GREEN,
            "detail": "Vienti / tuonti Cost Analysisista.",
            "blocker": "Laatu riippuu vaiheista 3–4.",
        },
        {
            "n": "6",
            "title": "Tuotantokäyttö",
            "status": "AVOIN",
            "fill": RED_LIGHT,
            "badge": RED,
            "detail": "Määrälaskentatiimi ilman päivittäistä kehittäjätukea.",
            "blocker": "Pilotti + hyväksymiskriteerit.",
        },
    ],
    "compact_labels": ["IFC", "Määrät", "TALO", "Littera", "Excel", "Tuot."],
    "h6": "6. Tuleva suunnitelma ja aikataulu (suuntaa-antava)",
    "crono_b_title": "Aikataulu B — Kohti tuotantokäyttöä (jaksot vahvistettava)",
    "crono_b_first": "Toiminto / kokonaisuus",
    "crono_b_periods": ["F1\nVk 1–4", "F2\nVk 4–8", "F3\nVk 8–12", "F4\nKk 3", "F5\nTuot."],
    "crono_b_rows": [
        ("Validoidut määrät (MVP-tyypit)", ["focus", "done", "done", "done", "done"]),
        ("TALO katettu (Revit + MCP asiakkaat)", ["partial", "focus", "done", "done", "done"]),
        ("Littera-säännöt (TALO → Littera)", ["empty", "partial", "focus", "done", "done"]),
        ("Pilotti Costan käyttäjillä (ilman päivittäistä tukea)", ["empty", "empty", "partial", "focus", "done"]),
        ("Tuotanto-MVP (tuotantokäyttö)", ["empty", "empty", "empty", "partial", "focus"]),
    ],
    "phase_headers": ["Vaihe", "Tavoite", "Tuotos", "Jakso (täydennä)"],
    "phase_rows": [
        ["F1", "Validoidut määrät sovituille IFC-tyypeille", "Sisäisesti hyväksytty Excel-referenssi", "[Vk 1–4]"],
        ["F2", "TALO katettu (sisäinen Revit + MCP asiakkaat)", "Menettely + 1 pilotti", "[Vk 4–8]"],
        ["F3", "Littera-säännöt (TALO → Littera)", "Sääntödokumentti + pilotti", "[Vk 8–12]"],
        ["F4", "Pilotti Costan käyttäjillä (ilman päivittäistä tukea)", "Go / no-go", "[Kk 3]"],
        ["F5", "Tuotanto-MVP (tuotantokäyttö)", "Tuotantopäivämäärä", "[Täydennä Jesse]"],
    ],
    "end": "— Luonnos v3 FI päättyy —",
}


def build_doc(content: dict) -> str:
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(1.8)
    section.bottom_margin = Cm(1.8)
    section.left_margin = Cm(2.0)
    section.right_margin = Cm(2.0)

    t = doc.add_paragraph()
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = t.add_run("Costa Laskenta Oy\n")
    r.bold = True
    r.font.size = Pt(14)
    r.font.color.rgb = RGBColor.from_string(BLUE)
    r2 = t.add_run("CostaBIM\n")
    r2.font.size = Pt(22)
    r2.bold = True
    r2.font.color.rgb = RGBColor.from_string(BLUE)
    r3 = t.add_run(content["cover_sub"] + "\n\n")
    r3.font.size = Pt(13)
    r3.bold = True
    r4 = t.add_run(f"{content['cover_note']} — {date.today().strftime('%d.%m.%Y')}\n")
    r4.font.size = Pt(10)
    r4.italic = True

    add_heading(doc, content["h1"], 1)
    add_para(doc, content["intro"])
    add_table(doc, content["q_headers"], content["q_rows"], col_widths=[5, 11])

    add_heading(doc, content["h2"], 1)
    add_para(doc, content["ctx"])
    add_table(doc, content["field_headers"], content["field_rows"], col_widths=[3.5, 4.5, 8])
    add_para(doc, content["talo_para"])

    add_heading(doc, content["schema1"], 2)
    add_flow_diagram(doc, content["flow1_title"], content["flow1"])
    add_para(doc, content["flow1_note"], size=10)

    add_heading(doc, content["h3"], 1)
    add_table(doc, content["via_headers"], content["via_rows"], col_widths=[4, 4.5, 3.5, 4])
    add_flow_diagram(doc, content["flow2_title"], content["flow2"])

    add_heading(doc, content["h4"], 1)
    t_a = add_cronograma(
        doc,
        content["crono_a_title"],
        content["crono_a_periods"],
        content["crono_a_rows"],
        PAST_HEADER,
    )
    set_first_crono_header(t_a, content["crono_a_first"])
    set_cell_shading(t_a.rows[0].cells[0], PAST_HEADER)
    add_legend(doc, content["legend"])
    add_para(doc, content["summary4"], size=10)

    add_heading(doc, content["h5"], 1)
    add_table(doc, content["gap_headers"], content["gap_rows"], col_widths=[5, 5, 6])

    doc.add_page_break()
    add_heading(doc, content["h_schema3"], 1)
    add_pipeline(
        doc,
        content["stages"],
        content["compact_labels"],
        [GREEN, GREEN, YELLOW, YELLOW, GREEN, RED],
        ["🟢", "🟢", "🟡", "🟡", "🟢", "🔴"],
        content["pipe_headers"],
    )

    add_heading(doc, content["h6"], 1)
    t_b = add_cronograma(
        doc,
        content["crono_b_title"],
        content["crono_b_periods"],
        content["crono_b_rows"],
        FUTURE_HEADER,
    )
    set_first_crono_header(t_b, content["crono_b_first"])
    set_cell_shading(t_b.rows[0].cells[0], FUTURE_HEADER)
    add_legend(doc, content["legend"])
    add_table(
        doc,
        content["phase_headers"],
        content["phase_rows"],
        col_widths=[1.2, 5.5, 5, 3.5],
        header_fill=FUTURE_HEADER,
    )

    add_para(doc, "\n" + content["end"], bold=True)

    out = os.path.join(DOWNLOADS, content["out"])
    doc.save(out)
    return out


def main():
    for pack in (ES, FI):
        path = build_doc(pack)
        print(f"Guardado: {path}")


if __name__ == "__main__":
    main()
