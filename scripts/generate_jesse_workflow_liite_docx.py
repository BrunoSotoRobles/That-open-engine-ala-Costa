#!/usr/bin/env python3
"""Generate Jesse workflow commentary + Liite — plain layout, minimal emphasis."""

from pathlib import Path

from docx import Document
from docx.shared import Pt, RGBColor

OUT = Path(r"C:\Users\BrunoSotoRobles\Downloads\CostaBIM_Claude_workflow_Liite.docx")
OUT_FALLBACK = Path(r"C:\Users\BrunoSotoRobles\Downloads\CostaBIM_Claude_workflow_Liite_uusi.docx")

# Light blue for CostaBIM-team additions (original Jesse text stays black).
ADDITION_COLOR = RGBColor(0x4A, 0x86, 0xC8)

SECTIONS = [
    {
        "original": [
            "1. Tutkitaan ja varmennetaan mallin käyttökelpoisuus",
            "a. omissa malleissa revit-nappi",
            "b. ulkopuolisissa ei selkeää tapaa etukäteen varmentaa käyttökelpoisuus.",
            "Kannattanee ajaa CB läpi ja C voi tutkia sen jälkeen ongelmia. Tähän on jo tehty työtä.",
        ],
        "tila": (
            "CB ajettu läpi → Excel + ModelHealth (korjaukset, abstain, varoitukset). "
            "C voi lukea yhteenvedon ja auditoida purun CB:n tuotoksista."
        ),
        "seuraava": "M1: selkokielinen yhteenveto määrälaskijalle: voit jatkaa / tarkista ensin X.",
        "ei_esta": (
            "Määräluettelo-luonnos voidaan tehdä, vaikka mallissa on tunnettuja puutteita — "
            "ne menevät kysymysraporttiin."
        ),
    },
    {
        "original": [
            "2. Päätetään objektit joista määrät otetaan. CB mahdollistaa eri rakenneosien jne valinnan.",
        ],
        "tila": "CB: Groups, IFC-tyypit, rajaus.",
        "seuraava": "M3: C kysyy scopen kerran alussa; määrälaskija vahvistaa.",
        "ei_esta": "Osa tyypeistä voi jäädä myöhemmäksi.",
    },
    {
        "original": [
            "3. Ajetaan malli CB läpi. Saadaan excel jossa kaikki valitut objektit listattuna kaikilla merkitsevillä ominaisuuksillaan.",
        ],
        "tila": "Send to Cost → Excel (IFC → Published → Costa-yksiköt) + ModelHealth.",
        "seuraava": "M1: vakiintunut paketti (Excel + ModelHealth samassa viennissä).",
        "ei_esta": "CB:n ydin — valmis pilottikäyttöön.",
    },
    {
        "original": [
            "4. Annetaan excel C:lle. C tekee määräluettelon Costan pohjaan.",
            "a. C voidaan käskeä kysymään kysymyksiä tarkentamaan määriä. Tässä ongelmana, että ilman vastausta C:n työt eivät etene joten järkevintä lienee joku raportti tms kysymyspatteristo jolla voidaan lopputuotosta tarkentaa sen jälkeen kun ensimmäinen tuotos on jo saatu. Tällöin myös joku muu voi jatkaa excelin tarkentamista Claudella.",
            "i. Tässä pointtina se, että muutamalla Costalaisella tulee olla kalliimpi lisenssi, mutta kaikille se ei ole tarpeen. Joten isoimman työn voisi tehdä kalliimmat lisenssit.",
        ],
        "tila": (
            "C saa Excelin CB:stä; Littera/TALO-työkalut käytössä. "
            "Yhtenäistä luonnos + kysymysraportti -pakettia ei vielä."
        ),
        "seuraava": "M2: luonnos aina ensin → kysymysraportti (Liite 2).",
        "ei_esta": (
            "C ei saa pysähtyä odottamaan vastauksia ennen ensimmäistä luonnosta. "
            "Kalliimpi C-lisenssi vain raskaimpaan työhön (kuten PDF:ssä)."
        ),
    },
    {
        "original": [
            "5. Juhon pelillä saadaan listoista (ovet, ikkunat, varusteet, kalusteet).",
        ],
        "tila": "Omat listauspohjat / CB-sarakkeet — rinnakkainen työ.",
        "seuraava": "M3: liitetään scopeen, kun pohja valmis kohteelle.",
        "ei_esta": "Yleinen määräluettelo-luonnos ei odota tätä.",
    },
    {
        "original": [
            "6. Rakennusselostus ja urakkarajaliite C:lle",
            "a. Täytyy varmaan määritellä mitä etsii urakkarajaliitteestä. Me etsitään tiettyjä urakan erityispiirteitä.",
            "b. Saisiko rakennusselostuksesta suoritteita tai rakenneosia määräluetteloon.",
        ],
        "tila": "Ei määritelty, mitä haetaan liitteestä.",
        "seuraava": "M5: lyhyt lista + 2–3 esimerkki-PDF.",
        "ei_esta": "IFC-pohjainen luonnos ilman näitä; rajaukset myöhemmin.",
    },
    {
        "original": [
            "7. C pystyy ”ennustamaan” mitä muita määriä/suoritteita kohteeseen tulee. Tätä pystytään varmaankin kouluttamaan jatkossa valmiilla määräluetteloilla eri rakennuskategorioista.",
            "a. C pystyy myös samalla prosessilla tarkistamaan määriä, että osuuko ne vaihteluväleihin. Tämä on mietinnän asteella.",
        ],
        "tila": "Ideataso.",
        "seuraava": "M5+: historialliset määräluettelot / säännöt.",
        "ei_esta": "—",
    },
    {
        "original": [
            "8. Tarkistetaan käsin samalla kun lasketaan puuttuvat määrät.",
            "a. Määrien tarkistus ei ole toivottua vaan tarkistusta tulee tehdä suoritteiden nimien perusteella ja suuruusluokka tarkastelulla.",
            "b. Suoritteiden nimiä täytyy päivittää, esim. betonilaadut, teräslaadut, elementtien nimet jne. C mahdollisesti saadaan päivittämään määräluetteloon suoritteiden nimiä esim. betonielementtien elementtityyppien mukaan ja lisäämään suoritteelle bruttoneliöt jne. Mutta ei ole tehty tämän eteen vielä töitä.",
        ],
        "tila": "Prosessisääntö; työkalutus kesken.",
        "seuraava": "M2/M4: tarkistus suoriten nimien ja suuruusluokan mukaan.",
        "ei_esta": "CB + C eivät korvaa määrälaskijaa.",
    },
    {
        "original": [
            "9. Hinnoittelu tulee ajaa C:lle vaikka kohde hinnoiteltaisiinkin käsin.",
            "a. C voi tehdä kahdesta laskelmasta vertailun ja mahdollisesti oppia lisää (oku?)",
        ],
        "tila": "Hinnoittelu CB:ssä olemassa; vertailuautomaatio ei.",
        "seuraava": "M5: esim. kahden laskelman vertailu.",
        "ei_esta": "Hinnoittelu voi olla käsin; C tuki.",
    },
]


def _run(paragraph, text: str, *, bold: bool = False, color: RGBColor | None = None) -> None:
    run = paragraph.add_run(text)
    run.bold = bold
    if color is not None:
        run.font.color.rgb = color


def original_line(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    _run(p, text)


def addition_line(doc: Document, text: str, *, bold: bool = False) -> None:
    p = doc.add_paragraph()
    _run(p, text, bold=bold, color=ADDITION_COLOR)


def addition_labeled(doc: Document, label: str, body: str) -> None:
    p = doc.add_paragraph()
    _run(p, label, bold=True, color=ADDITION_COLOR)
    _run(p, body, bold=False, color=ADDITION_COLOR)


def comment_block(doc: Document, tila: str, seuraava: str, ei_esta: str) -> None:
    addition_labeled(doc, "Tila nyt: ", tila)
    addition_labeled(doc, "Seuraava askel: ", seuraava)
    addition_labeled(doc, "Ei estä etenemistä: ", ei_esta)
    doc.add_paragraph()


def build_document() -> Document:
    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)

    # Jesse PDF opening
    original_line(doc, "CostaBim=CB, Claude =C")
    p_sub = doc.add_paragraph()
    _run(p_sub, "CB, C, määräluettelon ja hinnoittelun automatisointi", bold=True)
    doc.add_paragraph()

    addition_line(
        doc,
        "CB = CostaBIM (mallin katselu, määrädatan purku). C = Claude (määrälaskijan käyttöliittymä, auditointi). "
        "Claude on kytketty CostaBIM:iin. Määrälaskija hyväksyy suoritteet ja lopullisen määräluettelon.",
    )
    doc.add_paragraph()

    for section in SECTIONS:
        for line in section["original"]:
            original_line(doc, line)
        comment_block(doc, section["tila"], section["seuraava"], section["ei_esta"])

    original_line(
        doc,
        "Tässä enemmän tai vähemmän kaikki mahdolliset toimenpiteet jotka liittyvät CB, C, "
        "määräluettelo ja hinnoittelun automatisointiin. Riippuen kohteesta osa näistä jää pois.",
    )
    original_line(
        doc,
        "Rakennusliikkeelle tehtäviin toimenpiteisiin voi kuulua vielä esim. ennakoiden ryhmittely, "
        "vertailu jne jossa C voi auttaa huomattavasti.",
    )
    doc.add_paragraph()
    original_line(doc, "Riskeinä on")
    original_line(doc, "• ”nolojen” virheiden jääminen luetteloon (joita tulee kyllä käsinlaskettaessakin).")
    original_line(doc, "• laskelman näyttäminen liikaa tekoälyltä.")
    original_line(doc, "• laskijan laiskistuminen (ei tarkasteta tms kun määrä ja hinta kerta on)")
    original_line(doc, "• määrälaskijan ja CB:n ja C:n urakkaraja. Ettei mitään jää laskematta.")
    doc.add_paragraph()

    addition_line(doc, "Kommentti riskeihin:", bold=True)
    addition_line(doc, "ModelHealth + abstain dokumentoitu; määrälaskija tarkistaa suoritteet.")
    addition_line(doc, "Toimitus Costan tavallisessa muodossa — CB/C eivät näy asiakasmateriaalissa.")
    addition_line(doc, "C tuottaa luonnoksen ja kysymysraportin; määrälaskijan rooli eksplisiittinen.")
    addition_line(doc, "Scope ja kysymysraportti dokumentoivat urakkarajan.")
    doc.add_paragraph()
    doc.add_paragraph()

    addition_line(doc, "Liite 1: CB + C -workflow — nykytila, milestonet ja vastuut", bold=True)
    doc.add_paragraph()
    addition_line(
        doc,
        "Workflow perustuu edellä olevaan PDF:ään. CostaBIM (CB) purkaa mallin (Excel, ModelHealth). "
        "Claude (C) on määrälaskijan käyttöliittymä: luonnos, kysymysraportti, auditointi.",
    )
    doc.add_paragraph()
    addition_line(doc, "Ei-estävä sääntö:", bold=True)
    addition_line(doc, "1. Ensimmäinen määräluettelo-luonnos tulee aina.")
    addition_line(doc, "2. Tarkentavat kysymykset erillisessä raportissa (Liite 2).")
    addition_line(doc, "3. Tarkistus suoritteiden nimien ja suuruusluokan mukaan.")
    doc.add_paragraph()
    addition_line(doc, "Milestonet:", bold=True)
    addition_line(doc, "M0. Nykytila — pilottiin valmis (CB: Excel, ModelHealth; C: ohjaa CB:tä).")
    addition_line(doc, "M1. Luotettava purku + yhteenveto määrälaskijalle.")
    addition_line(doc, "M2. Määräluettelo-luonnos + kysymysraportti (ei blokkaa).")
    addition_line(doc, "M3. Scope + Juho-listat tarpeen mukaan.")
    addition_line(doc, "M4. Pilotti oikealla kohteella (määrälaskija).")
    addition_line(
        doc,
        "M5+. Urakkarajaliite, rakennusselostus, ennuste, hinnoitteluvertailu — ei estä M1–M4.",
    )
    doc.add_paragraph()
    addition_line(doc, "Tiivistelmä:", bold=True)
    addition_line(
        doc,
        "Määrälaskija työskentelee Clauden kanssa. Claude on kytketty CostaBIM:iin. "
        "CostaBIM purkaa mallin. Claude tekee luonnoksen ja kysymysraportin. Määrälaskija tarkistaa suoritteet.",
    )
    doc.add_paragraph()
    doc.add_paragraph()

    addition_line(doc, "Liite 2: Kysymysraportin pohja", bold=True)
    doc.add_paragraph()
    addition_line(doc, "Kysymysraportti — [kohde], [päivä]")
    addition_line(doc, "1. CB abstain / ei korjannut — mitä ja miksi (ModelHealth)")
    addition_line(doc, "2. Puuttuvat TALO / Littera")
    addition_line(doc, "3. Tarkista suoritteen nimi")
    addition_line(doc, "4. Suuruusluokka — poikkeavat arvot")
    addition_line(doc, "5. Scope ulkopuolella")
    addition_line(doc, "6. Seuraava toimenpide — määrälaskijan päätös")

    return doc


def main() -> None:
    doc = build_document()
    OUT.parent.mkdir(parents=True, exist_ok=True)
    target = OUT
    try:
        doc.save(target)
    except PermissionError:
        target = OUT_FALLBACK
        doc.save(target)
        print(f"Note: {OUT.name} locked — wrote fallback instead.")
    print(f"Wrote: {target}")


if __name__ == "__main__":
    main()
