"""Patch Ville's mallinnusohje docx with CostaBim §3.5 and §4.4 content."""
from __future__ import annotations

from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph


SRC = Path(
    r"C:\Users\BrunoSotoRobles\Downloads\Mallinnusohjeet arkkitehdille ja rakennesuunnittelijalle.docx"
)
DST = Path(
    r"C:\Users\BrunoSotoRobles\Downloads\Mallinnusohjeet arkkitehdille ja rakennesuunnittelijalle (CostaBim).docx"
)

PLACEHOLDER = "Täydennetään myöhemmin"

ARK_BLOCKS: list[tuple[str, str]] = [
    (
        "normal",
        "ARK-mallin tulee täyttää Costa Laskennan esitarkistus: IFC-rakenne ja mittatiedot ovat "
        "kunnossa ennen määrä- ja litteratyötä. Lopullinen määrä- ja litteratyö tehdään Costa "
        "Laskennan työprosessissa mallin pohjalta.",
    ),
    ("heading", "Tilamalli (IfcSpace)"),
    ("bullet", "Jokaisella tilalla nimi, tilatyyppi ja osasto (kohta 3.1)."),
    (
        "bullet",
        "Huoneseloste: pinnat, sisäovet, sisäikkunat, kalusteet ja varusteet mallinnetaan "
        "malliin — ei pelkästään piirustuksissa.",
    ),
    ("heading", "Mitattavat objektit"),
    (
        "bullet",
        "Ovet, ikkunat, lattiat, alakatot, kalusteet: IFC-mittatiedot mukana (Qto / "
        "suunnitteluohjelman määrätiedot).",
    ),
    (
        "bullet",
        "Rakenteet (kohta 3.3): nimeäminen pysyy vakaana rakennetyypin mukaan (väliseinät, "
        "ulkoseinät, vesikatot jne.).",
    ),
    ("heading", "Vesikatot (IfcRoof / IfcSlab)"),
    (
        "bullet",
        "Jos vesikatto viedään IfcSlab-tyyppinä, projektin sisäinen käytäntö on yhtenäinen ja "
        "dokumentoitu muutoslistassa.",
    ),
    (
        "bullet",
        "Vesikattoa ja kantavaa laattaa ei sekoiteta ilman selkeää erottelua nimessä tai "
        "tyypissä.",
    ),
    ("heading", "Kerrosrakenteet (sandwich) — IfcWall, IfcSlab, IfcRoof"),
    (
        "bullet",
        "Jos rakennustuote viedään useana IFC-elementtinä (esim. kuorielementti + eriste + "
        "sisäkuori; kerroslaatta; katto kerroksittain), nimien on oltava johdonmukaisia koko "
        "projektissa.",
    ),
    (
        "bullet",
        "Pääelementti (kuori / kantava osa) ja kerrokset on erotettava selkeästi nimillä.",
    ),
    (
        "bullet",
        "Costa tunnistaa pääelementin ja kerrokset mallin vastaanotossa, ryhmittelee ne ennen "
        "määrälaskentaa, eikä kerroksia lasketa erillisinä paikkoina samasta tuotteesta.",
    ),
    ("heading", "Toimitus"),
    ("bullet", "Päivitetty IFC + muutoslista jokaisessa toimituksessa (kohta 2)."),
    (
        "normal",
        "Malli, joka täyttää nämä vaatimukset, helpottaa määränottoa CostaBIMissä. Se ei korvaa "
        "Costa Laskennan tarkistusta eikä takaa valmista hintalaskelmaa ilman työprosessia.",
    ),
]

RAK_BLOCKS: list[tuple[str, str]] = [
    (
        "normal",
        "RAK-mallin tulee täyttää Costa Laskennan esitarkistus: IFC-rakenne ja mittatiedot ovat "
        "kunnossa ennen määrä- ja litteratyötä. Lopullinen määrä- ja litteratyö tehdään Costa "
        "Laskennan työprosessissa mallin pohjalta.",
    ),
    ("heading", "IFC-mittatiedot"),
    (
        "bullet",
        "Mitattavilla objekteilla Tekla Quantity -tiedot (Length, Width, Height, Volume, Net "
        "surface area jne.) tai vastaavat BaseQuantities / Qto-tiedot.",
    ),
    ("bullet", "Ei tuplaobjekteja; nimeäminen pysyy vakaana koko projektin ajan (kohta 4.3)."),
    (
        "bullet",
        "Nimet ovat johdonmukaisia IfcWall-, IfcBeam-, IfcColumn-, IfcSlab- ja IfcFooting "
        "-elementeissä.",
    ),
    ("heading", "Kerrosrakenteet (sandwich) — IfcWall, IfcSlab, IfcRoof"),
    (
        "bullet",
        "Jos rakennustuote viedään useana IFC-elementtinä (esim. julkisivun kuori + eriste + "
        "verhous; kerros- tai ontelolaatta; katto kerroksittain), nimien on oltava "
        "johdonmukaisia koko projektissa.",
    ),
    (
        "bullet",
        "Pääelementti ja kerrokset on erotettava selkeästi nimillä (esim. KUORIELEMENTTI vs. "
        "ERISTE / SISÄKUORI tai projektin vastaava käytäntö).",
    ),
    (
        "bullet",
        "Costa tunnistaa pääelementin ja kerrokset mallin vastaanotossa, ryhmittelee ne ennen "
        "määrälaskentaa, eikä kerroksia lasketa erillisinä paikkoina samasta tuotteesta.",
    ),
    ("heading", "Toimitus"),
    ("bullet", "Päivitetty IFC + muutoslista jokaisessa toimituksessa (kohta 2)."),
    (
        "normal",
        "Malli, joka täyttää nämä vaatimukset, helpottaa määränottoa CostaBIMissä. Se ei korvaa "
        "Costa Laskennan tarkistusta eikä takaa valmista hintalaskelmaa ilman työprosessia.",
    ),
]


def _clone_paragraph_after(reference: Paragraph, text: str, style_name: str) -> Paragraph:
    new_p = deepcopy(reference._element)
    reference._element.addnext(new_p)
    paragraph = Paragraph(new_p, reference._parent)
    paragraph.style = style_name
    paragraph.text = text
    return paragraph


def _style_paragraph(paragraph: Paragraph, kind: str) -> None:
    if kind == "heading":
        paragraph.style = "Normal"
        if paragraph.runs:
            paragraph.runs[0].bold = True
        else:
            run = paragraph.add_run(paragraph.text)
            paragraph.text = ""
            run.bold = True
    elif kind == "bullet":
        paragraph.style = "List Paragraph"
        if not paragraph.text.startswith("•"):
            paragraph.text = f"• {paragraph.text.lstrip('•').strip()}"


def replace_placeholder(doc: Document, section_marker: str, blocks: list[tuple[str, str]]) -> None:
    placeholder_idx = None
    for i, paragraph in enumerate(doc.paragraphs):
        if PLACEHOLDER in paragraph.text and section_marker in doc.paragraphs[i - 1].text:
            placeholder_idx = i
            break
    if placeholder_idx is None:
        raise RuntimeError(f"Placeholder not found for section {section_marker!r}")

    anchor = doc.paragraphs[placeholder_idx]
    first_kind, first_text = blocks[0]
    anchor.text = first_text
    anchor.style = "Normal"
    _style_paragraph(anchor, first_kind)

    insert_before = doc.paragraphs[placeholder_idx + 1]._element
    for kind, text in blocks[1:]:
        new_el = deepcopy(anchor._element)
        insert_before.addprevious(new_el)
        paragraph = Paragraph(new_el, anchor._parent)
        paragraph.text = text
        _style_paragraph(paragraph, kind)


def main() -> None:
    if not SRC.is_file():
        raise SystemExit(f"Source not found: {SRC}")

    doc = Document(str(SRC))
    replace_placeholder(doc, "3.5", ARK_BLOCKS)
    replace_placeholder(doc, "4.4", RAK_BLOCKS)
    doc.save(str(DST))
    print(f"Saved: {DST}")


if __name__ == "__main__":
    main()
